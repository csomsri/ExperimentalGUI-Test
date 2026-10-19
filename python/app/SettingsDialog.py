# -*- coding: utf-8 -*-
"""
SettingsDialog.py — unified Settings dialog (RF & Vacuum, Alarm, Plot/Logging, Beam calibration)

Changes vs older builds:
- No manual/digital selector (mode is auto-driven by LabVIEW nibble in main_control_gui.py)
- Beam calibration editor is curve-only (no linear mode)
"""

from __future__ import annotations
import os
from PyQt6.QtWidgets import (
    QDialog, QTabWidget, QWidget, QVBoxLayout, QGridLayout, QLabel, QDoubleSpinBox,
    QLineEdit, QCheckBox, QPushButton, QHBoxLayout, QComboBox, QSpinBox, QPlainTextEdit,
    QFileDialog, QMessageBox
)
from PyQt6.QtCore import Qt

from app_settings import app_settings, save as save_app
from alarm import alarm_manager, AlarmConfig
from beam_cal import beam_cal, save_cfg as save_beam_cfg

# Optional Excel support for LUT import
try:
    from openpyxl import load_workbook
    _HAVE_XLSX = True
except Exception:
    _HAVE_XLSX = False


class SettingsDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Settings")
        self.resize(760, 560)

        # ---- Dark style ----
        self.setStyleSheet("""
            QDialog { background-color: rgba(14,16,24,235); }
            QTabWidget::pane {
                border: 1px solid #1C5F82;
                padding: 8px;
                margin-top: -2px;
                background: #0E1421;
                border-radius: 6px;
            }
            QTabBar::tab {
                background: #0F1A2A;
                color: #AEE1FF;
                border: 1px solid #2D6B93;
                border-bottom-color: #1C5F82;
                padding: 6px 12px;
                margin-right: 2px;
                border-top-left-radius: 6px;
                border-top-right-radius: 6px;
            }
            QTabBar::tab:hover   { background: #123252; }
            QTabBar::tab:selected{
                background: #1A324B; color: #EAF6FF;
                border: 1px solid #3C8BC0;
                border-bottom-color: #1A324B;
            }
            QLabel  { color:#EAF6FF; font-size:12px; }
            QLineEdit,QDoubleSpinBox,QComboBox,QSpinBox,QPlainTextEdit {
                background:#0E1421; color:#EAF6FF;
                border:1px solid #2D6B93; border-radius:4px; padding:4px 6px;
                selection-background-color:#2D6B93;
            }
            QPlainTextEdit { font-family: Consolas, Menlo, Monospace; font-size:12px; }
            QCheckBox { color:#EAF6FF; }
            QPushButton {
                background:#1F74C1; color:white; border:1px solid #2C88DA;
                border-radius:6px; padding:6px 12px;
            }
            QPushButton:hover  { background:#2896F0; }
            QPushButton:pressed{ background:#1A5FA1; }
        """)

        # Tabs
        tabs = QTabWidget(self)
        t_rfvac, t_alarm, t_plot, t_beam = QWidget(), QWidget(), QWidget(), QWidget()
        tabs.addTab(t_rfvac, "RF & Vacuum")
        tabs.addTab(t_alarm, "Alarm")
        tabs.addTab(t_plot,  "Plot/Logging")
        tabs.addTab(t_beam,  "Beam")

        # ---------------------- RF & Vacuum ----------------------
        rv = QGridLayout(t_rfvac)
        s = app_settings

        row = 0
        rv.addWidget(self._h("RF display & filter"), row, 0, 1, 4); row += 1
        self.rf_max = QDoubleSpinBox(); self.rf_max.setRange(1, 10000); self.rf_max.setDecimals(1); self.rf_max.setValue(s.rf_gauge_max_kv)
        self.rf_db  = QDoubleSpinBox(); self.rf_db.setRange(0, 1000);  self.rf_db.setDecimals(2); self.rf_db.setValue(s.rf_deadband_kv)
        self.rf_tau = QDoubleSpinBox(); self.rf_tau.setRange(0, 10);   self.rf_tau.setDecimals(2); self.rf_tau.setValue(s.rf_smooth_tau_s)
        rv.addWidget(QLabel("Gauge full-scale (kV)"), row, 0, Qt.AlignmentFlag.AlignRight); rv.addWidget(self.rf_max, row, 1)
        rv.addWidget(QLabel("Deadband (kV)"),        row, 2, Qt.AlignmentFlag.AlignRight); rv.addWidget(self.rf_db,  row, 3); row += 1
        rv.addWidget(QLabel("Smoothing τ (s)"),      row, 0, Qt.AlignmentFlag.AlignRight); rv.addWidget(self.rf_tau, row, 1); row += 1

        row += 1
        rv.addWidget(self._h("Vacuum display & filter"), row, 0, 1, 4); row += 1
        self.vac_units = QComboBox(); self.vac_units.addItems(["mbar", "torr"]); self.vac_units.setCurrentText(s.vac_units)
        self.vac_min   = QDoubleSpinBox(); self.vac_min.setRange(-12, -1); self.vac_min.setDecimals(0); self.vac_min.setValue(s.vac_min_exp)
        self.vac_max   = QDoubleSpinBox(); self.vac_max.setRange(-12, -1); self.vac_max.setDecimals(0); self.vac_max.setValue(s.vac_max_exp)
        self.vac_tau   = QDoubleSpinBox(); self.vac_tau.setRange(0, 10);   self.vac_tau.setDecimals(2); self.vac_tau.setValue(s.vac_smooth_tau_s)
        rv.addWidget(QLabel("Units"),              row, 0, Qt.AlignmentFlag.AlignRight); rv.addWidget(self.vac_units, row, 1)
        rv.addWidget(QLabel("Gauge decades (min)"),row, 2, Qt.AlignmentFlag.AlignRight); rv.addWidget(self.vac_min,   row, 3); row += 1
        rv.addWidget(QLabel("Gauge decades (max)"),row, 2, Qt.AlignmentFlag.AlignRight); rv.addWidget(self.vac_max,   row, 3); row += 1
        rv.addWidget(QLabel("Smoothing τ (s)"),    row, 0, Qt.AlignmentFlag.AlignRight); rv.addWidget(self.vac_tau,   row, 1); row += 1

        # Optional: list which vac channels are TP gauges
        self.chk_tp = [QCheckBox(n) for n in ("Vac1","Vac2","Vac3","Vac4","Vac5")]
        if hasattr(s, "vac_tp_keys"):
            for i,k in enumerate(["vac1","vac2","vac3","vac4","vac5"]):
                if k in s.vac_tp_keys: self.chk_tp[i].setChecked(True)
        rv.addWidget(QLabel("TP gauges:"), row, 0, Qt.AlignmentFlag.AlignRight)
        rv.addWidget(self.chk_tp[0], row,1); rv.addWidget(self.chk_tp[1], row,2); rv.addWidget(self.chk_tp[2], row,3); row += 1
        rv.addWidget(self.chk_tp[3], row,1); rv.addWidget(self.chk_tp[4], row,2)

        # ---------------------- Alarm ----------------------
        al = QGridLayout(t_alarm)
        ac = alarm_manager.cfg

        al.addWidget(self._h("RF alarm"), 0,0,1,4)
        self.rf_dkv = QDoubleSpinBox(); self.rf_dkv.setRange(0,1e6);  self.rf_dkv.setDecimals(2); self.rf_dkv.setValue(ac.rf_dkv)
        self.rf_win = QDoubleSpinBox(); self.rf_win.setRange(0.1,60); self.rf_win.setDecimals(2); self.rf_win.setValue(ac.rf_window_s)
        self.rf_bg  = QLineEdit(ac.rf_bg); self.rf_bg.setPlaceholderText("rf_alarm_bg.png")
        al.addWidget(QLabel("ΔRF (kV) ≥"), 1,0, Qt.AlignmentFlag.AlignRight); al.addWidget(self.rf_dkv, 1,1)
        al.addWidget(QLabel("within (s)"), 1,2, Qt.AlignmentFlag.AlignRight); al.addWidget(self.rf_win, 1,3)
        al.addWidget(QLabel("RF overlay image"), 2,0, Qt.AlignmentFlag.AlignRight); al.addWidget(self.rf_bg, 2,1,1,3)

        al.addWidget(self._h("Vacuum alarm (rise by factor)"), 3,0,1,4)
        self.vac_fac = QDoubleSpinBox(); self.vac_fac.setRange(1.0,1e6); self.vac_fac.setDecimals(2); self.vac_fac.setValue(ac.vac_factor)
        self.vac_win = QDoubleSpinBox(); self.vac_win.setRange(0.1,120); self.vac_win.setDecimals(2); self.vac_win.setValue(ac.vac_window_s)
        self.vac_bg  = QLineEdit(ac.vac_bg); self.vac_bg.setPlaceholderText("vac_alarm_bg.png")
        al.addWidget(QLabel("Factor ≥"), 4,0, Qt.AlignmentFlag.AlignRight); al.addWidget(self.vac_fac, 4,1)
        al.addWidget(QLabel("within (s)"), 4,2, Qt.AlignmentFlag.AlignRight); al.addWidget(self.vac_win, 4,3)
        al.addWidget(QLabel("Vac overlay image"), 5,0, Qt.AlignmentFlag.AlignRight); al.addWidget(self.vac_bg, 5,1,1,3)

        self.chk_watch = [QCheckBox(n) for n in ("Vac1","Vac2","Vac3","Vac4","Vac5")]
        for i,k in enumerate(["vac1","vac2","vac3","vac4","vac5"]):
            if k in ac.vac_channels: self.chk_watch[i].setChecked(True)
        al.addWidget(QLabel("Watch:"), 6,0, Qt.AlignmentFlag.AlignRight)
        al.addWidget(self.chk_watch[0], 6,1); al.addWidget(self.chk_watch[1], 6,2); al.addWidget(self.chk_watch[2], 6,3)
        al.addWidget(self.chk_watch[3], 7,1); al.addWidget(self.chk_watch[4], 7,2)

        # ---------------------- Plot/Logging ----------------------
        pl = QGridLayout(t_plot)
        self.plot_fps = QDoubleSpinBox(); self.plot_fps.setRange(1,240);   self.plot_fps.setDecimals(0); self.plot_fps.setValue(s.plot_fps)
        self.plot_sps = QDoubleSpinBox(); self.plot_sps.setRange(10,5000); self.plot_sps.setDecimals(0); self.plot_sps.setValue(s.plot_sps)
        self.plot_win = QDoubleSpinBox(); self.plot_win.setRange(1,3600);  self.plot_win.setDecimals(0); self.plot_win.setValue(s.plot_window_s)
        self.log_cap  = QDoubleSpinBox(); self.log_cap.setRange(0,1e6);    self.log_cap.setDecimals(0)
        self.log_cap.setValue(0 if s.log_cap_hz in (None,0) else float(s.log_cap_hz))
        self.log_batch= QDoubleSpinBox(); self.log_batch.setRange(1,50000); self.log_batch.setDecimals(0); self.log_batch.setValue(s.log_batch)
        self.log_flush= QDoubleSpinBox(); self.log_flush.setRange(10,5000); self.log_flush.setDecimals(0); self.log_flush.setValue(s.log_flush_ms)

        pl.addWidget(QLabel("Plot FPS"),        0,0, Qt.AlignmentFlag.AlignRight); pl.addWidget(self.plot_fps, 0,1)
        pl.addWidget(QLabel("Display SPS"),     1,0, Qt.AlignmentFlag.AlignRight); pl.addWidget(self.plot_sps, 1,1)
        pl.addWidget(QLabel("Time window (s)"), 2,0, Qt.AlignmentFlag.AlignRight); pl.addWidget(self.plot_win, 2,1)
        pl.addWidget(QLabel("Log cap (Hz) 0=∞"),3,0, Qt.AlignmentFlag.AlignRight); pl.addWidget(self.log_cap,  3,1)
        pl.addWidget(QLabel("Log batch rows"),  4,0, Qt.AlignmentFlag.AlignRight); pl.addWidget(self.log_batch,4,1)
        pl.addWidget(QLabel("Log flush (ms)"),  5,0, Qt.AlignmentFlag.AlignRight); pl.addWidget(self.log_flush,5,1)

        # ---------------------- Beam (curve-only) ----------------------
        b = QGridLayout(t_beam)
        bc = beam_cal.cfg

        b.addWidget(self._h("Range selection (auto by LabVIEW)"), 0,0,1,4)
        self.man_idx = QSpinBox(); self.man_idx.setRange(0, max(0, len(bc.ranges)-1)); self.man_idx.setValue(bc.manual_index)
        self.dig_src = QComboBox(); self.dig_src.addItems(["beam_range_idx","bitmask_low4","bitmask_high4"])
        self.dig_src.setCurrentText(bc.digital_source)
        b.addWidget(QLabel("Manual index"), 1,0, Qt.AlignmentFlag.AlignRight); b.addWidget(self.man_idx, 1,1)
        b.addWidget(QLabel("Digital src"),  1,2, Qt.AlignmentFlag.AlignRight); b.addWidget(self.dig_src, 1,3)

        b.addWidget(self._h("Edit range (curve only)"), 2,0,1,4)
        self.range_sel = QComboBox()
        for i,r in enumerate(bc.ranges):
            self.range_sel.addItem(f"#{i}  {r.name}")
        self.range_sel.setCurrentIndex(bc.manual_index if 0 <= bc.manual_index < len(bc.ranges) else 0)
        b.addWidget(QLabel("Range"), 3,0, Qt.AlignmentFlag.AlignRight); b.addWidget(self.range_sel, 3,1)

        self.r_name   = QLineEdit()
        self.r_fs_nA  = QDoubleSpinBox(); self.r_fs_nA.setRange(1e-6, 1e9); self.r_fs_nA.setDecimals(6)
        self.r_mode   = QComboBox(); self.r_mode.addItems(["curve"])
        self.r_gain   = QDoubleSpinBox(); self.r_gain.setRange(-1e12,1e12); self.r_gain.setDecimals(12)
        self.r_offs   = QDoubleSpinBox(); self.r_offs.setRange(-1e12,1e12); self.r_offs.setDecimals(6)
        b.addWidget(QLabel("Name"),       4,0, Qt.AlignmentFlag.AlignRight); b.addWidget(self.r_name, 4,1)
        b.addWidget(QLabel("FS (nA)"),    4,2, Qt.AlignmentFlag.AlignRight); b.addWidget(self.r_fs_nA, 4,3)
        b.addWidget(QLabel("Gain (nA/V)"),5,0, Qt.AlignmentFlag.AlignRight); b.addWidget(self.r_gain, 5,1)
        b.addWidget(QLabel("Offset (nA)"),5,2, Qt.AlignmentFlag.AlignRight); b.addWidget(self.r_offs, 5,3)

        b.addWidget(self._h("Curve (LUT) — lines 'V, I'"), 6,0,1,4)
        self.curve_text = QPlainTextEdit()
        self.curve_text.setPlaceholderText("e.g.\n0.339497681, 1.0\n0.606, 2.0\n...")
        b.addWidget(self.curve_text, 7,0,1,4)

        # Buttons: Apply curve + Load from Excel
        btn_row = QHBoxLayout()
        self.btn_load_excel = QPushButton("Load curve from Excel…")
        self.btn_load_excel.clicked.connect(self._on_load_curve_excel)
        btn_curve = QPushButton("Apply curve to this range")
        btn_curve.clicked.connect(self._on_apply_curve)
        btn_row.addStretch(1); btn_row.addWidget(self.btn_load_excel); btn_row.addWidget(btn_curve)
        b.addLayout(btn_row, 8, 0, 1, 4)

        # Load selected range into editor fields
        self.range_sel.currentIndexChanged.connect(self._load_range_into_editor)
        self._load_range_into_editor(self.range_sel.currentIndex())

        # ---- Footer ----
        footer = QHBoxLayout()
        footer.addStretch(1)
        btn_ok = QPushButton("Save");   btn_ok.clicked.connect(self._on_save)
        btn_ca = QPushButton("Cancel"); btn_ca.clicked.connect(self.reject)
        footer.addWidget(btn_ok); footer.addWidget(btn_ca)

        out = QVBoxLayout(self)
        out.addWidget(tabs)
        out.addLayout(footer)

    # ---------- Helpers ----------
    def _h(self, txt: str) -> QLabel:
        l = QLabel(txt); l.setStyleSheet("color:#AEE1FF; font-weight:bold;")
        return l

    # ---------- Beam tab helpers ----------
    def _load_range_into_editor(self, idx: int):
        bc = beam_cal.cfg
        if not (0 <= idx < len(bc.ranges)): return
        r = bc.ranges[idx]
        self.r_name.setText(r.name)
        self.r_fs_nA.setValue(float(r.fs_nA))
        self.r_mode.setCurrentText("curve")
        self.r_gain.setValue(float(r.gain_na_per_v))
        self.r_offs.setValue(float(r.offset_na))
        # show curve in text box (if any)
        self.curve_text.clear()
        if r.curve and r.mode == "curve":
            for v,i in sorted(r.curve):
                self.curve_text.appendPlainText(f"{v:.9f}, {i:.9f}")

    def _store_editor_into_range(self, idx: int):
        bc = beam_cal.cfg
        if not (0 <= idx < len(bc.ranges)): return
        r = bc.ranges[idx]
        r.name = self.r_name.text().strip() or r.name
        r.fs_nA = float(self.r_fs_nA.value())
        r.mode  = "curve"  # locked to curve
        r.gain_na_per_v = float(self.r_gain.value())
        r.offset_na     = float(self.r_offs.value())
        bc.ranges[idx] = r

    # Unit-aware parser for the text box (supports: "nA", "µA/uA", "pA")
    def _parse_curve_text_to_na(self, txt: str):
        """
        Accepts lines like:
            1.700000000, 0.100000000
            1.740000000  0.200000000
            1.770000000, 1.0 nA
            1.820000000, 250 pA
            1.900000000, 0.5 µA
        Returns sorted list of (V, I_nA).
        """
        pts = []
        for line in (txt or "").splitlines():
            s = line.strip()
            if not s or s.startswith("#"):
                continue
            parts = [p for p in s.replace(",", " ").split() if p]
            if len(parts) < 2:
                continue
            v = float(parts[0])

            raw  = parts[1]
            unit = parts[2] if len(parts) >= 3 else ""
            try:
                i = float(raw)
            except ValueError:
                # e.g. "0.3nA" or "300pA"
                num = ''.join(ch for ch in raw if (ch.isdigit() or ch in ".-+eE"))
                suf = raw[len(num):].strip()
                i = float(num)
                unit = unit or suf

            u = unit.lower().replace("μ","u")
            if   u.startswith("pa"): i_na = i / 1000.0     # pA -> nA
            elif u.startswith("ua"): i_na = i * 1000.0     # µA -> nA
            else:                    i_na = i              # default nA
            pts.append((v, i_na))

        # de-dup by V, sort ascending by V
        dedup = {}
        for v,i in pts:
            dedup[float(v)] = float(i)
        return sorted((float(v), float(i)) for v,i in dedup.items())

    # Apply curve from text area
    def _on_apply_curve(self):
        idx = self.range_sel.currentIndex()
        self._store_editor_into_range(idx)  # keep name/fs/gain/offset edits
        pts = self._parse_curve_text_to_na(self.curve_text.toPlainText())
        if len(pts) < 2:
            return
        r = beam_cal.cfg.ranges[idx]
        r.mode  = "curve"
        r.curve = pts
        self.r_mode.setCurrentText("curve")
        if len(pts) <= 6:
            print(f"[BeamCal] Range#{idx} curve points:", pts)
        else:
            print(f"[BeamCal] Range#{idx} curve points: {len(pts)} (V∈[{pts[0][0]:.6f},{pts[-1][0]:.6f}])")

    # Load curve from Excel file (two columns: V, I — I may have units)
    def _on_load_curve_excel(self):
        if not _HAVE_XLSX:
            QMessageBox.warning(self, "Excel not available",
                                "openpyxl is not installed. Install it or paste the curve into the text box.")
            return
        path, _ = QFileDialog.getOpenFileName(self, "Select Excel curve", "",
                                              "Excel (*.xlsx *.xls);;All Files (*)")
        if not path:
            return
        try:
            wb = load_workbook(path, data_only=True)
            ws = wb.active
            lines = []
            # Assume header on first row; start at row 2
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or (row[0] is None) or (row[1] is None):
                    continue
                v = str(row[0]).strip()
                i = str(row[1]).strip()
                if v == "" or i == "":
                    continue
                lines.append(f"{v}, {i}")
            if len(lines) < 2:
                QMessageBox.warning(self, "Invalid data",
                                    "Need at least 2 numeric rows with V and I columns.")
                return
            self.curve_text.clear()
            self.curve_text.appendPlainText("\n".join(lines))
            self._on_apply_curve()
            QMessageBox.information(self, "Curve loaded",
                                    f"Loaded {len(lines)} points from:\n{os.path.basename(path)}")
        except Exception as e:
            QMessageBox.critical(self, "Read error", f"Failed to read Excel file:\n{e}")

    # ---------- Save ----------
    def _on_save(self):
        # RF/Vac
        s = app_settings
        s.rf_gauge_max_kv  = float(self.rf_max.value())
        s.rf_deadband_kv   = float(self.rf_db.value())
        s.rf_smooth_tau_s  = float(self.rf_tau.value())
        s.vac_units        = self.vac_units.currentText()
        s.vac_min_exp      = int(self.vac_min.value())
        s.vac_max_exp      = int(self.vac_max.value())
        s.vac_smooth_tau_s = float(self.vac_tau.value())
        save_app(s)

        # Alarm
        vac_watch = []
        for i,k in enumerate(["vac1","vac2","vac3","vac4","vac5"]):
            if self.chk_watch[i].isChecked(): vac_watch.append(k)
        if not vac_watch: vac_watch = ["vac1","vac2"]
        cfg = AlarmConfig(
            rf_dkv=float(self.rf_dkv.value()),
            rf_window_s=float(self.rf_win.value()),
            vac_factor=float(self.vac_fac.value()),
            vac_window_s=float(self.vac_win.value()),
            vac_channels=vac_watch,
            rf_bg=self.rf_bg.text().strip(),
            vac_bg=self.vac_bg.text().strip(),
        )
        alarm_manager.apply_config(cfg)

        # Plot/Logging
        cap = float(self.log_cap.value())
        s.plot_fps      = float(self.plot_fps.value())
        s.plot_sps      = float(self.plot_sps.value())
        s.plot_window_s = float(self.plot_win.value())
        s.log_cap_hz    = None if cap <= 0 else cap
        s.log_batch     = int(self.log_batch.value())
        s.log_flush_ms  = int(self.log_flush.value())
        save_app(s)

        # Beam (selection metadata only — mode is auto in main_control_gui)
        bc = beam_cal.cfg
        bc.manual_index   = int(self.man_idx.value())
        bc.digital_source = self.dig_src.currentText()

        # Persist the current range edits
        idx = self.range_sel.currentIndex()
        self._store_editor_into_range(idx)
        save_beam_cfg(bc)

        self.accept()

