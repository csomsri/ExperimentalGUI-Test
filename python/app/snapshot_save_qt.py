from project_paths import project_file
# snapshot_save_qt.py
# Snapshot Saver (Qt): note dialog + write snapshot row to Excel.
# - Creates the Excel file if it doesn't exist
# - Auto-expands headers if new keys appear later
# - Stores: timestamp (human), ts (epoch), note, + all signal keys
# - Prevents collisions if incoming signals contain keys like "timestamp", "ts", "note"
# - Stores SCALED values (engineering units) using scaling.get_scaler() if available
# - NEW: fills "holes" (deleted rows) by writing to the first truly empty row

import os
import time
from datetime import datetime
from typing import Dict, Any, List, Optional

from PyQt6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QPushButton
)
from PyQt6.QtCore import Qt

try:
    from openpyxl import Workbook, load_workbook
    _HAVE_OPENPYXL = True
except Exception:
    _HAVE_OPENPYXL = False

# Try to import your global scaler (no hard dependency)
try:
    from scaling import get_scaler
    _HAVE_SCALER = True
except Exception:
    _HAVE_SCALER = False
    get_scaler = None


DARK_STYLE = """
QDialog, QWidget {
    background-color: #0b0f14;
    color: #e6f3ff;
    font-family: "Segoe UI", sans-serif;
    font-size: 12px;
}
QLabel { color: #e6f3ff; }
QLineEdit {
    background-color: #111821;
    color: #e6f3ff;
    border: 1px solid #2b92f9;
    border-radius: 8px;
    padding: 6px 8px;
}
QPushButton {
    background-color: #0e5aa6;
    color: #ffffff;
    border: 1px solid #67c2ff;
    border-radius: 10px;
    padding: 8px 12px;
}
QPushButton:hover { background-color: #1299c8; }
QPushButton:disabled {
    background-color: #1a2a3a;
    color: #7ea4c7;
    border-color: #2b3e54;
}
"""


def _to_excel_safe(v: Any) -> Any:
    """Convert values into something Excel/openpyxl can store safely."""
    if v is None:
        return None
    if isinstance(v, bool):
        return bool(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:
        s = str(v).strip()
        if s == "":
            return None
        return float(s)
    except Exception:
        return str(v)


def _scale_signals_if_possible(signals: Dict[str, Any]) -> Dict[str, Any]:
    """
    Scale signals using the global scaler if available.

    Handles BOTH:
      - scalar keys (e.g., 'ch1', 'beam_current', 'vac1')
      - packed arrays (e.g., 'channels' list holding [ch1..ch12, main_magnet, centering_beam])
    """
    if not _HAVE_SCALER:
        return dict(signals)

    try:
        scaler = get_scaler()
    except Exception:
        return dict(signals)

    if scaler is None or not hasattr(scaler, "data") or not hasattr(scaler, "raw_to_eng"):
        return dict(signals)

    # Import canonical key order for packed arrays
    try:
        from scaling import TRIM_KEYS, EXTR_BASE_KEYS, EXTR_ANGLE_KEYS, SOURCE_KEYS, TRANSPORT_KEYS
    except Exception:
        TRIM_KEYS = []
        EXTR_BASE_KEYS = []
        EXTR_ANGLE_KEYS = []
        SOURCE_KEYS = []
        TRANSPORT_KEYS = []

    raw = dict(signals)
    out: Dict[str, Any] = {}

    def scale_scalar(key: str, val):
        try:
            if key in scaler.data:
                return scaler.raw_to_eng(key, val)
        except Exception:
            pass
        return val

    def scale_list(keys: List[str], vals):
        if not isinstance(vals, (list, tuple)):
            return vals
        if len(keys) == 0 or len(vals) != len(keys):
            return vals
        scaled = []
        for k, v in zip(keys, vals):
            scaled.append(scale_scalar(k, v))
        return scaled

    for k, v in raw.items():
        # --- packed arrays ---
        if k == "channels":
            out[k] = scale_list(TRIM_KEYS, v)  # [ch1..ch12, main_magnet, centering_beam]
            continue

        if k in ("extraction", "extraction_base"):
            out[k] = scale_list(EXTR_BASE_KEYS, v)
            continue

        if k in ("extraction_angles", "angles"):
            out[k] = scale_list(EXTR_ANGLE_KEYS, v)
            continue

        if k in ("source", "source_pack"):
            out[k] = scale_list(SOURCE_KEYS, v)
            continue

        if k in ("transport", "transport_pack"):
            out[k] = scale_list(TRANSPORT_KEYS, v)
            continue

        # --- normal scalar channels ---
        out[k] = scale_scalar(k, v)

    return out


def _is_empty_cell(val: Any) -> bool:
    return val is None or (isinstance(val, str) and val.strip() == "")


def _find_first_empty_row(ws, ncols: int) -> int:
    """
    Finds the first truly empty row by checking ALL columns 1..ncols.
    Starts at row 2 (row 1 is header).
    Includes one-past-max_row so it can append at the end if no holes.
    """
    # scan from row 2 downward, include one past max_row
    for r in range(2, ws.max_row + 2):
        any_filled = False
        for c in range(1, ncols + 1):
            v = ws.cell(row=r, column=c).value
            if not _is_empty_cell(v):
                any_filled = True
                break
        if not any_filled:
            return r
    return ws.max_row + 1


def append_snapshot_xlsx(
    filepath: str,
    signals: Dict[str, Any],
    note: str = "",
    ts_epoch: Optional[float] = None,
    timestamp_str: Optional[str] = None,
    sheet_name: str = str(project_file('snapshots')),
    scaled_only: bool = True,
) -> str:
    """
    Append one snapshot row to an Excel file.

    Columns:
      timestamp (human), ts (epoch), note, <signal keys...>

    If scaled_only=True:
      tries to scale signals using scaling.get_scaler().

    Key collision rules:
      - If signals contains 'timestamp', it is renamed to 'timestamp_src'
      - If signals contains 'ts', it is renamed to 'ts_src'
      - If signals contains 'note', it is renamed to 'note_src'

    NEW behavior:
      - writes into the first truly empty row (fills holes after deletions)
    """
    if not _HAVE_OPENPYXL:
        raise RuntimeError("openpyxl is not installed. Install it with: pip install openpyxl")

    if not isinstance(signals, dict) or not signals:
        raise ValueError("signals must be a non-empty dict")

    # Apply scaling (no need to change main_control_gui.py)
    sig_in = _scale_signals_if_possible(signals) if scaled_only else dict(signals)

    # Ensure directory exists
    folder = os.path.dirname(os.path.abspath(filepath))
    if folder and not os.path.exists(folder):
        os.makedirs(folder, exist_ok=True)

    ts_epoch = float(ts_epoch if ts_epoch is not None else time.time())
    timestamp_str = timestamp_str or datetime.fromtimestamp(ts_epoch).strftime("%Y-%m-%d %H:%M:%S")

    # --- Prevent collisions with our metadata columns ---
    sig = dict(sig_in)
    if "timestamp" in sig:
        sig["timestamp_src"] = sig.pop("timestamp")
    if "ts" in sig:
        sig["ts_src"] = sig.pop("ts")
    if "note" in sig:
        sig["note_src"] = sig.pop("note")

    base_cols = ["timestamp", "ts", "note"]

    # Normalize for Excel
    norm_signals: Dict[str, Any] = {str(k): _to_excel_safe(v) for k, v in sig.items()}
    signal_cols = sorted(norm_signals.keys())

    # ---------- Create file if missing ----------
    if not os.path.exists(filepath):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        headers = base_cols + signal_cols
        ws.append(headers)

        values_map: Dict[str, Any] = {
            "timestamp": timestamp_str,
            "ts": ts_epoch,
            "note": note,
            **norm_signals,
        }
        row = [values_map.get(h, None) for h in headers]

        # Fill first empty row (will be row 2 in a new file)
        row_idx = _find_first_empty_row(ws, ncols=len(headers))
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx).value = val

        wb.save(filepath)
        return os.path.abspath(filepath)

    # ---------- Open existing file ----------
    wb = load_workbook(filepath)
    ws = wb.active
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

    # Read existing headers
    existing_headers: List[str] = []
    for cell in ws[1]:
        existing_headers.append(str(cell.value) if cell.value is not None else "")

    # Ensure base columns exist
    for col in base_cols:
        if col not in existing_headers:
            existing_headers.append(col)
            ws.cell(row=1, column=len(existing_headers)).value = col

    # Ensure all signal columns exist
    for col in signal_cols:
        if col not in existing_headers:
            existing_headers.append(col)
            ws.cell(row=1, column=len(existing_headers)).value = col

    # Build row aligned to headers
    values_map2: Dict[str, Any] = {
        "timestamp": timestamp_str,
        "ts": ts_epoch,
        "note": note,
        **norm_signals,
    }
    row2 = [values_map2.get(h, None) for h in existing_headers]

    # NEW: fill first empty row instead of always appending at bottom
    row_idx2 = _find_first_empty_row(ws, ncols=len(existing_headers))
    for col_idx, val in enumerate(row2, start=1):
        ws.cell(row=row_idx2, column=col_idx).value = val

    wb.save(filepath)
    return os.path.abspath(filepath)


class SnapshotSaveDialog(QDialog):
    """Dialog shown when Snapshot is clicked: shows file path + asks for note."""
    def __init__(self, parent=None, filepath: str = str(project_file('snapshot_log.xlsx')), default_note: str = ""):
        super().__init__(parent)
        self.setWindowTitle("Save Snapshot")
        self.setModal(True)
        self.resize(560, 170)

        self.filepath = filepath
        self.note = default_note

        root = QVBoxLayout(self)

        # File
        row1 = QHBoxLayout()
        row1.addWidget(QLabel("File:"))
        self.file_edit = QLineEdit(filepath)
        self.file_edit.setReadOnly(True)
        self.file_edit.setCursorPosition(0)
        row1.addWidget(self.file_edit, 1)
        root.addLayout(row1)

        # Note
        row2 = QHBoxLayout()
        row2.addWidget(QLabel("Note:"))
        self.note_edit = QLineEdit(default_note)
        self.note_edit.setPlaceholderText("e.g., 'Stable beam @ 20 MeV', 'After TC10 tweak', ...")
        row2.addWidget(self.note_edit, 1)
        root.addLayout(row2)

        # Buttons
        row3 = QHBoxLayout()
        row3.addStretch(1)
        self.btn_save = QPushButton("Save")
        self.btn_cancel = QPushButton("Cancel")
        self.btn_save.clicked.connect(self._on_save)
        self.btn_cancel.clicked.connect(self.reject)
        row3.addWidget(self.btn_save)
        row3.addWidget(self.btn_cancel)
        root.addLayout(row3)

        self.setStyleSheet(DARK_STYLE)
        self.note_edit.setFocus(Qt.FocusReason.ActiveWindowFocusReason)

    def _on_save(self):
        self.note = (self.note_edit.text() or "").strip()
        self.accept()
