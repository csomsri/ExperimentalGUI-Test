# snapshot_recall_qt.py
# Snapshot Recall — Preview ALL (4 panels) + Apply per-panel buttons
#
# Expected snapshot Excel format (row-based snapshots):
# - columns include: timestamp (or timestamp_human), note
# - list-valued columns like: channels, source, extraction, extraction_angles, transport
#   (may be actual lists or stringified lists like "[0.1, 0.2, ...]")
#
# Behavior:
# - Previews all panels for the selected snapshot
# - Magnetic preview shows: TC1..TC12 + Main Magnet + Centering Beam
# - Apply Magnetic applies ONLY TC1..TC12 via apply_callback (safe)
# - Apply Source/Extraction/Transport call optional callbacks if provided; otherwise show "not wired yet"

import os
import ast
from datetime import datetime
from typing import Dict, List, Optional, Callable, Any

from PyQt6.QtWidgets import (
    QDialog, QWidget, QVBoxLayout, QHBoxLayout, QLabel, QComboBox,
    QTableWidget, QTableWidgetItem, QMessageBox, QPushButton,
    QGridLayout, QFrame, QSizePolicy
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QPainter, QColor, QPen, QLinearGradient, QFont

from openpyxl import load_workbook


# -----------------------------
# Helpers
# -----------------------------
def _parse_list(val) -> List[float]:
    """Accept list/tuple or stringified list; return list of floats when possible."""
    if val is None:
        return []
    if isinstance(val, (list, tuple)):
        out = []
        for x in val:
            try:
                out.append(float(x))
            except Exception:
                out.append(0.0)
        return out

    if isinstance(val, str):
        s = val.strip()
        if not s:
            return []
        try:
            obj = ast.literal_eval(s)
            if isinstance(obj, (list, tuple)):
                out = []
                for x in obj:
                    try:
                        out.append(float(x))
                    except Exception:
                        out.append(0.0)
                return out
        except Exception:
            return []

    return []


def _parse_dt(val) -> Optional[datetime]:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    s = str(val).strip().replace("Z", "")
    if not s:
        return None

    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M",
    ):
        try:
            return datetime.strptime(s, fmt)
        except Exception:
            pass

    try:
        return datetime.fromisoformat(s)
    except Exception:
        return None


def _first_existing(headers: List[str], candidates: List[str]) -> Optional[str]:
    hmap = {h.lower(): h for h in headers}
    for c in candidates:
        if c.lower() in hmap:
            return hmap[c.lower()]
    return None


TS_CANDIDATES   = ["timestamp", "timestamp_human", "time", "ts", "saved_at", "datetime"]
NOTE_CANDIDATES = ["note", "title", "label", "comment", "reference"]


# -----------------------------
# Labels (pretty names)
# -----------------------------
MAG_LABELS = [f"TC{i}" for i in range(1, 13)] + ["Main Magnet", "Centering Beam"]

SRC_LABELS = ["Filament A", "Filament V", "Arc V", "Arc A", "Gas Flow", "Gas Turns"]
EXT_LABELS_MAIN = ["Coax A", "Inside Iron A", "Outside Iron A", "ESD V", "ESD A", "Fringe A"]
EXT_LABELS_ANG  = ["ESD In °", "ESD Out °", "Coax In °", "Coax Out °", "Iron In °", "Iron Out °"]
TRN_LABELS = [
    "Q1A A", "Q1B A", "Q2A A", "Q2B A",
    "Steer40 A", "Steer50 A",
    "Coll S1 L A", "Coll S1 R A", "Coll S2 L A", "Coll S2 R A",
]


# -----------------------------
# UI: Neon border-only panel
# -----------------------------
class NeonPanel(QFrame):
    """Draws only the neon border. Title is handled by layout header widgets (no overlap)."""
    def __init__(self):
        super().__init__()
        self.setStyleSheet("background: rgba(10,14,22,210);")
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)

    def paintEvent(self, ev):
        super().paintEvent(ev)
        p = QPainter(self)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)

        r = self.rect().adjusted(6, 6, -6, -6)

        grad = QLinearGradient(r.left(), r.center().y(), r.right(), r.center().y())
        grad.setColorAt(0.0, QColor(0, 220, 255))
        grad.setColorAt(0.55, QColor(160, 120, 255))
        grad.setColorAt(1.0, QColor(255, 90, 200))

        p.setPen(QPen(grad, 3))
        p.setBrush(Qt.BrushStyle.NoBrush)
        p.drawRoundedRect(r, 12, 12)


# -----------------------------
# Callback types
# -----------------------------
ApplyTC_Callback   = Callable[[Dict[str, float]], None]   # TC1..TC12 payload: {"ch1":..., ...}
ApplyList_Callback = Callable[[List[float]], None]        # generic list payload


# -----------------------------
# Recall dialog
# -----------------------------
class SnapshotRecallDialog(QDialog):
    def __init__(
        self,
        parent=None,
        snapshot_path: str = "",
        apply_callback: Optional[ApplyTC_Callback] = None,     # magnetic TC-only
        apply_source: Optional[ApplyList_Callback] = None,
        apply_extraction: Optional[ApplyList_Callback] = None,
        apply_transport: Optional[ApplyList_Callback] = None,
    ):
        super().__init__(parent)
        self.snapshot_path = snapshot_path

        # Callbacks
        self.apply_tc = apply_callback
        self.apply_source = apply_source
        self.apply_extraction = apply_extraction
        self.apply_transport = apply_transport

        self.headers: List[str] = []
        self.rows: List[Dict[str, Any]] = []

        self.setWindowTitle("Snapshot Recall — Preview All, Apply Per Panel")
        self.resize(1600, 950)
        self.setWindowFlags(self.windowFlags() | Qt.WindowType.WindowMaximizeButtonHint)

        root = QVBoxLayout(self)

        # Top bar
        top = QHBoxLayout()

        self.combo = QComboBox()
        self.combo.setMinimumWidth(900)
        self.combo.currentIndexChanged.connect(self._on_select)

        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.reject)

        top.addWidget(QLabel("Snapshot:"))
        top.addWidget(self.combo)
        top.addStretch(1)
        top.addWidget(btn_close)
        root.addLayout(top)

        self.info = QLabel("")
        root.addWidget(self.info)

        # Panels (2x2)
        grid = QGridLayout()
        root.addLayout(grid, 1)

        self.mag = self._make_panel("Magnetic Field (TC1–TC12 + Main + Centering)", "Apply Magnetic", self._apply_tc_only)
        self.src = self._make_panel("Beam Source", "Apply Source", self._apply_source_only)
        self.ext = self._make_panel("Beam Extraction", "Apply Extraction", self._apply_extraction_only)
        self.trn = self._make_panel("Beam Transport", "Apply Transport", self._apply_transport_only)

        grid.addWidget(self.mag["panel"], 0, 0)
        grid.addWidget(self.src["panel"], 0, 1)
        grid.addWidget(self.ext["panel"], 1, 0)
        grid.addWidget(self.trn["panel"], 1, 1)

        grid.setRowStretch(0, 1)
        grid.setRowStretch(1, 1)
        grid.setColumnStretch(0, 1)
        grid.setColumnStretch(1, 1)

        self._load_file()

    def _make_panel(self, title: str, apply_text: str, apply_fn) -> Dict[str, Any]:
        panel = NeonPanel()
        lay = QVBoxLayout(panel)
        lay.setContentsMargins(14, 14, 14, 14)
        lay.setSpacing(8)

        # Header row (layout-managed: no overlap)
        header = QHBoxLayout()
        header.setContentsMargins(6, 4, 6, 4)

        lbl = QLabel(title)
        lbl.setFont(QFont("Segoe UI", 11, QFont.Weight.DemiBold))
        lbl.setStyleSheet("color: rgb(220,245,255);")
        header.addWidget(lbl)
        header.addStretch(1)

        btn_apply = QPushButton(apply_text)
        btn_apply.setFixedHeight(32)
        btn_apply.setMinimumWidth(150)
        btn_apply.clicked.connect(apply_fn)
        btn_apply.setStyleSheet("""
            QPushButton {
                background: rgba(15,35,55,220);
                border: 1px solid rgb(90,180,255);
                border-radius: 8px;
                padding: 4px 14px;
                color: rgb(220,245,255);
            }
            QPushButton:hover {
                background: rgba(30,60,90,230);
                border-color: rgb(160,220,255);
            }
            QPushButton:pressed {
                background: rgba(10,25,40,240);
            }
            QPushButton:disabled {
                background: rgba(30,30,30,120);
                color: rgba(200,200,200,120);
                border: 1px solid rgba(120,120,120,90);
            }
        """)
        header.addWidget(btn_apply)

        lay.addLayout(header)

        # Table
        table = QTableWidget(0, 2)
        table.setHorizontalHeaderLabels(["Signal", "Value"])
        table.horizontalHeader().setStretchLastSection(True)
        table.verticalHeader().setVisible(False)
        table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        table.setSelectionMode(QTableWidget.SelectionMode.NoSelection)
        table.setStyleSheet("""
            QTableWidget {
                background: rgba(5,10,18,200);
                gridline-color: rgba(80,120,160,80);
                border: 1px solid rgba(80,140,200,120);
                border-radius: 8px;
            }
            QHeaderView::section {
                background: rgba(20,40,65,220);
                color: rgb(210,240,255);
                padding: 6px;
                border: none;
            }
        """)
        lay.addWidget(table, 1)

        return {"panel": panel, "table": table, "btn": btn_apply}

    # -------------------------
    # LOAD EXCEL (headers fixed)
    # -------------------------
    def _load_file(self):
        if not self.snapshot_path or not os.path.exists(self.snapshot_path):
            self.info.setText(f"Snapshot file not found:\n{self.snapshot_path}")
            self._set_apply_enabled(False)
            return

        try:
            wb = load_workbook(self.snapshot_path, data_only=True, read_only=True)
            ws = wb.active

            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            self.headers = [str(h).strip() if h is not None else "" for h in header_row]

            self.rows = []
            for r in ws.iter_rows(min_row=2, values_only=True):
                if r and any(v is not None for v in r):
                    self.rows.append(dict(zip(self.headers, r)))

        except Exception as e:
            self.info.setText(f"Failed to read snapshot file:\n{self.snapshot_path}\n\n{e}")
            self._set_apply_enabled(False)
            return

        ts_col = _first_existing(self.headers, TS_CANDIDATES)
        note_col = _first_existing(self.headers, NOTE_CANDIDATES)

        self.combo.blockSignals(True)
        self.combo.clear()

        for i, row in enumerate(reversed(self.rows)):  # newest first
            ts_val = row.get(ts_col) if ts_col else None
            note_val = row.get(note_col) if note_col else None

            dt = _parse_dt(ts_val)
            ts_txt = dt.strftime("%Y-%m-%d %H:%M:%S") if dt else "no-time"
            note_txt = str(note_val).strip() if note_val is not None else ""

            label = f"{ts_txt} — {note_txt}" if note_txt else ts_txt
            actual_index = len(self.rows) - 1 - i
            self.combo.addItem(label, actual_index)

        self.combo.blockSignals(False)

        self.info.setText(f"Loaded {len(self.rows)} snapshots from:\n{os.path.abspath(self.snapshot_path)}")
        self._set_apply_enabled(bool(self.rows))

        if self.rows:
            self._on_select(0)

    def _set_apply_enabled(self, enabled: bool):
        self.mag["btn"].setEnabled(enabled)
        self.src["btn"].setEnabled(enabled)
        self.ext["btn"].setEnabled(enabled)
        self.trn["btn"].setEnabled(enabled)

    # -------------------------
    # SELECT SNAPSHOT → FILL TABLES
    # -------------------------
    def _on_select(self, _):
        idx = self.combo.currentData()
        if idx is None or not self.rows:
            return

        row = self.rows[int(idx)]

        channels = _parse_list(row.get("channels"))
        source = _parse_list(row.get("source"))
        extraction = _parse_list(row.get("extraction"))
        angles = _parse_list(row.get("extraction_angles"))
        transport = _parse_list(row.get("transport"))

        # Magnetic: TC1..TC12 + main + centering
        mag_vals = channels[:14]  # first 14 if present
        while len(mag_vals) < 14:
            mag_vals.append(0.0)
        self._fill_table(self.mag["table"], MAG_LABELS, mag_vals[:14])

        # Source
        self._fill_table(self.src["table"], SRC_LABELS[:len(source)], source)

        # Extraction (main + angles)
        ext_vals = extraction[:len(EXT_LABELS_MAIN)]
        ang_vals = angles[:len(EXT_LABELS_ANG)]
        self._fill_table(
            self.ext["table"],
            EXT_LABELS_MAIN[:len(ext_vals)] + EXT_LABELS_ANG[:len(ang_vals)],
            ext_vals + ang_vals
        )

        # Transport
        self._fill_table(self.trn["table"], TRN_LABELS[:len(transport)], transport)

    def _fill_table(self, table: QTableWidget, labels: List[str], values: List[float]):
        table.setRowCount(0)
        for name, val in zip(labels, values):
            r = table.rowCount()
            table.insertRow(r)

            it_name = QTableWidgetItem(str(name))
            it_val = QTableWidgetItem(f"{float(val):.6g}")

            it_name.setFlags(it_name.flags() & ~Qt.ItemFlag.ItemIsEditable)
            it_val.setFlags(it_val.flags() & ~Qt.ItemFlag.ItemIsEditable)
            it_val.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)

            table.setItem(r, 0, it_name)
            table.setItem(r, 1, it_val)

    # -------------------------
    # APPLY buttons (per panel)
    # -------------------------
    def _apply_tc_only(self):
        idx = self.combo.currentData()
        if idx is None or self.apply_tc is None:
            QMessageBox.information(self, "Apply Magnetic", "Magnetic apply is not wired.")
            return

        row = self.rows[int(idx)]
        channels = _parse_list(row.get("channels"))

        # Apply ONLY TC1..TC12 (ignore main/centering)
        payload: Dict[str, float] = {}
        for i in range(12):
            payload[f"ch{i+1}"] = float(channels[i]) if i < len(channels) else 0.0

        res = QMessageBox.question(
            self,
            "Apply Magnetic (TC-only)?",
            "Apply ONLY TC1–TC12 targets from this snapshot?\n\n"
            "Main Magnet and Centering Beam will NOT be applied.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if res != QMessageBox.StandardButton.Yes:
            return

        self.apply_tc(payload)
        QMessageBox.information(self, "Apply Magnetic", "Applied TC1–TC12.")

    def _apply_source_only(self):
        idx = self.combo.currentData()
        if idx is None:
            return
        row = self.rows[int(idx)]
        source = _parse_list(row.get("source"))

        if self.apply_source is None:
            QMessageBox.information(self, "Apply Source", "Source apply is not wired yet.")
            return

        if QMessageBox.question(
            self, "Apply Source?", "Apply Source settings from this snapshot?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) != QMessageBox.StandardButton.Yes:
            return

        self.apply_source(source)
        QMessageBox.information(self, "Apply Source", "Source settings applied.")

    def _apply_extraction_only(self):
        idx = self.combo.currentData()
        if idx is None:
            return
        row = self.rows[int(idx)]
        extraction = _parse_list(row.get("extraction"))
        angles = _parse_list(row.get("extraction_angles"))
        payload = extraction + angles

        if self.apply_extraction is None:
            QMessageBox.information(self, "Apply Extraction", "Extraction apply is not wired yet.")
            return

        if QMessageBox.question(
            self, "Apply Extraction?", "Apply Extraction settings from this snapshot?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) != QMessageBox.StandardButton.Yes:
            return

        self.apply_extraction(payload)
        QMessageBox.information(self, "Apply Extraction", "Extraction settings applied.")

    def _apply_transport_only(self):
        idx = self.combo.currentData()
        if idx is None:
            return
        row = self.rows[int(idx)]
        transport = _parse_list(row.get("transport"))

        if self.apply_transport is None:
            QMessageBox.information(self, "Apply Transport", "Transport apply is not wired yet.")
            return

        if QMessageBox.question(
            self, "Apply Transport?", "Apply Transport settings from this snapshot?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) != QMessageBox.StandardButton.Yes:
            return

        self.apply_transport(transport)
        QMessageBox.information(self, "Apply Transport", "Transport settings applied.")
