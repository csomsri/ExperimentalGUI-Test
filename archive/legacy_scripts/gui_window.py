# -*- coding: utf-8 -*-
"""
Main GUI with SQLite viewer + Snapshot-to-Excel
"""

import sys, os, sqlite3, csv
from datetime import datetime
from PyQt6.QtWidgets import (
    QApplication, QWidget, QPushButton, QLabel, QLineEdit, QMessageBox, QFileDialog
)
from PyQt6.QtGui import QFont, QIcon, QPixmap, QPainter
from PyQt6.QtCore import Qt, QSize

from sqlite_viewer_qt import SQLiteLogViewer
from MagneticFieldMonitoringWindow import MagneticFieldMonitoringWindow
from MagneticFieldControllerWindow import MagneticFieldControllerWindow

# ---------- Config (adjust these if needed) ----------
DEFAULT_DB_PATH = "cyclotron_data.db"        # e.g., r"C:\path\to\cyclotron_data.db"
TABLE_NAME      = "channel_data"
TIME_COL        = "timestamp"
EXPECTED_COLS   = [
    "ch1","ch2","ch3","ch4","ch5","ch6","ch7","ch8","ch9","ch10","ch11","ch12",
    "main_magnet","centering_beam"
]
SNAPSHOT_XLSX   = "trimcoil_snapshots.xlsx"  # fallback to CSV if openpyxl isn't available
SNAPSHOT_CSV    = "trimcoil_snapshots.csv"
SNAPSHOT_ICON   = "icon_snapshot.png"        # provide this file
FALLBACK_ICON   = "icon_db.png"              # used if snapshot icon missing

# Excel writer (optional)
try:
    from openpyxl import Workbook, load_workbook
    _HAVE_OPENPYXL = True
except Exception:
    _HAVE_OPENPYXL = False


def _load_icon(path: str, fallback: str | None = None) -> QIcon:
    """Return a QIcon if file exists, else fallback/blank."""
    if path and os.path.exists(path):
        return QIcon(path)
    if fallback and os.path.exists(fallback):
        return QIcon(fallback)
    return QIcon()


def _parse_db_time(val) -> datetime:
    """
    Convert DB time (epoch or text) -> datetime (local).
    Accepts: numeric epoch (int/float/str), ISO 'YYYY-MM-DD HH:MM:SS[.us]', 'YYYY-MM-DDTHH:MM:SS'.
    """
    # numeric epoch?
    try:
        return datetime.fromtimestamp(float(val))
    except Exception:
        pass

    s = str(val).replace("Z", "")
    # try common layouts
    for fmt in ("%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M:%S.%f",
                "%Y-%m-%dT%H:%M:%S",
                "%Y-%m-%dT%H:%M:%S.%f",
                "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt)
        except Exception:
            continue
    # last resort: fromisoformat (covers many variants)
    return datetime.fromisoformat(s)


class ControllerWindow(QWidget):
    def __init__(self, title):
        super().__init__()
        self.setWindowTitle(title)
        self.setGeometry(300, 300, 400, 200)
        self.setStyleSheet("background-color: #222; color: white;")
        label = QLabel(f"{title} - Controller Window", self)
        label.setFont(QFont("Arial", 14))
        label.setGeometry(50, 80, 300, 40)
        label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.show()


class MainControlWindow(QWidget):
    def __init__(self, data_queue=None, target_values=None, control_queue=None):
        super().__init__()
        self.data_queue = data_queue
        self.target_values = target_values
        self.control_queue = control_queue

        self.mag_monitor = None
        self.field_controller = None
        self._generic_windows = []  # keep refs to generic controller windows
        self.db_path = DEFAULT_DB_PATH

        self.setWindowTitle("Cyclotron Control System")
        self.setGeometry(100, 100, 1152, 768)

        self.background = QPixmap("crocker_ui_background.jpg")
        if self.background.isNull():
            print("⚠️ Background image not found.")

        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        self.setStyleSheet("background: transparent;")

        # Note input for Snapshot
        self.note_edit = QLineEdit(self)
        self.note_edit.setPlaceholderText("Record note / name")
        self.note_edit.setStyleSheet(
            "background:#111; color:white; border:1px solid #66c2ff; border-radius:6px; padding:6px;"
        )

        self.init_buttons()

    # ---------- Windows ----------
    def show_magnetic_field_monitoring(self):
        if self.mag_monitor is None or not self.mag_monitor.isVisible():
            self.mag_monitor = MagneticFieldMonitoringWindow(self.data_queue, self.target_values)
            self.mag_monitor.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose)
            self.mag_monitor.show()
        else:
            self.mag_monitor.raise_()
            self.mag_monitor.activateWindow()

    def show_magnetic_field_controller(self):
        if self.field_controller is None or not self.field_controller.isVisible():
            self.field_controller = MagneticFieldControllerWindow(self.control_queue)
            self.field_controller.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose)
            self.field_controller.show()
        else:
            self.field_controller.raise_()
            self.field_controller.activateWindow()

    def show_sqlite_viewer(self):
        try:
            if not hasattr(self, "_sqlite_viewer") or self._sqlite_viewer is None or not self._sqlite_viewer.isVisible():
                self._sqlite_viewer = SQLiteLogViewer()
                self._sqlite_viewer.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose)
                self._sqlite_viewer.show()
            else:
                self._sqlite_viewer.raise_()
                self._sqlite_viewer.activateWindow()
        except Exception as e:
            print(f"SQLite viewer error: {e}")

    def open_generic_controller(self, title: str):
        w = ControllerWindow(title)
        w.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose)
        self._generic_windows.append(w)
        w.show()

    # ---------- Painting ----------
    def paintEvent(self, event):
        painter = QPainter(self)
        scaled = self.background.scaled(
            self.size(),
            Qt.AspectRatioMode.KeepAspectRatioByExpanding,
            Qt.TransformationMode.SmoothTransformation,
        )
        painter.drawPixmap(self.rect(), scaled)

    # ---------- Buttons + Snapshot ----------
    def init_buttons(self):
        button_font = QFont("Segoe UI", 9)

        self.button_config = [
            {"label": "Magnetic Field", "tooltip": "Magnetic Field Monitoring", "x": 953, "y": 104, "size": 138, "icon": "icon_magnetic.png"},
            {"label": "Source",         "tooltip": "Extraction + Source Monitoring", "x": 1035, "y": 260, "size": 115, "icon": "icon_extraction.png"},
            {"label": "Transport",      "tooltip": "Beam Transport Monitoring", "x": 987, "y": 400, "size": 110, "icon": "icon_transport.png"},
            {"label": "RF Power",       "tooltip": "RF Power Monitoring", "x": 888, "y": 510, "size": 110, "icon": "icon_rf.png"},
            {"label": "Field Ctrl",     "tooltip": "Magnetic Field Controller", "x": 461, "y": 104, "size": 138, "icon": "icon_magnetic_ctrl.png"},
            {"label": "Source Ctrl",    "tooltip": "Extraction + Source Controller", "x": 400, "y": 260, "size": 115, "icon": "icon_extraction_ctrl.png"},
            {"label": "Transport Ctrl", "tooltip": "Beam Transport Controller", "x": 445, "y": 400, "size": 110, "icon": "icon_transport_ctrl.png"},
            {"label": "RF Ctrl",        "tooltip": "RF Power Controller", "x": 540, "y": 510, "size": 110, "icon": "icon_rf_ctrl.png"},

            {"label": "SQLite",         "tooltip": "SQLite Log Viewer", "x": 1100, "y": 320, "size": 110, "icon": "icon_db.png"},

            # New Snapshot button with icon
            {"label": "Snapshot",       "tooltip": "Save current channels to Excel", "x": 200, "y": 455, "size": 110, "icon": SNAPSHOT_ICON},
        ]

        snapshot_btn_geom = None

        for cfg in self.button_config:
            btn = QPushButton("", self)
            btn.setToolTip(cfg["tooltip"])

            # Load icon with fallback for Snapshot
            icon_path = cfg["icon"]
            if cfg["label"] == "Snapshot":
                btn.setIcon(_load_icon(icon_path, FALLBACK_ICON))
            else:
                btn.setIcon(_load_icon(icon_path))
            icon_size = int(cfg["size"] * 0.6)
            btn.setIconSize(QSize(icon_size, icon_size))

            btn.setGeometry(cfg["x"], cfg["y"], cfg["size"], cfg["size"])
            btn.setStyleSheet(f"""
                QPushButton {{
                    background-color: #007ACC;
                    border: 2px solid #66c2ff;
                    border-radius: {cfg["size"] // 2}px;
                }}
                QPushButton:hover {{ background-color: #00bcd4; }}
            """)

            tip = cfg["tooltip"]
            if tip == "Magnetic Field Controller":
                btn.clicked.connect(self.show_magnetic_field_controller)
            elif tip == "Magnetic Field Monitoring":
                btn.clicked.connect(self.show_magnetic_field_monitoring)
            elif tip == "SQLite Log Viewer":
                btn.clicked.connect(self.show_sqlite_viewer)
            elif tip == "Save current channels to Excel":
                btn.clicked.connect(self.snapshot_to_excel)
                snapshot_btn_geom = (cfg["x"], cfg["y"], cfg["size"])
            elif tip.endswith("Controller"):
                btn.clicked.connect(lambda _, t=tip: self.open_generic_controller(t))
            else:
                btn.clicked.connect(lambda _, t=tip: print(f"Monitoring: {t}"))

            label = QLabel(cfg["label"], self)
            label.setGeometry(cfg["x"], cfg["y"] + cfg["size"] + 5, cfg["size"], 20)
            label.setFont(button_font)
            label.setStyleSheet("color: white;")
            label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        # Place note box near Snapshot button
        if snapshot_btn_geom:
            x, y, size = snapshot_btn_geom
            self.note_edit.setGeometry(x - 240, y + size//2 - 14, 220, 28)

    # ---------- Snapshot internals ----------
    def _ensure_db_path(self) -> str | None:
        path = self.db_path
        if path and os.path.exists(path):
            return path
        path, _ = QFileDialog.getOpenFileName(self, "Locate cyclotron SQLite DB", "", "SQLite DB (*.db *.sqlite *.sqlite3);;All Files (*)")
        if path:
            self.db_path = path
            return path
        return None

    def _get_latest_row(self, con: sqlite3.Connection):
        cur = con.cursor()
        cols = [r[1] for r in cur.execute(f"PRAGMA table_info({TABLE_NAME})").fetchall()]
        if TIME_COL not in cols:
            raise RuntimeError(f"Time column '{TIME_COL}' not found in '{TABLE_NAME}'.")
        present = [c for c in EXPECTED_COLS if c in cols]
        if not present:
            raise RuntimeError("No expected channel columns found.")
        fields = [TIME_COL] + present
        row = cur.execute(f"SELECT {', '.join(fields)} FROM {TABLE_NAME} ORDER BY ROWID DESC LIMIT 1").fetchone()
        if not row:
            raise RuntimeError("No rows found in the table.")
        return fields, row

    def _append_to_xlsx(self, fields, row_values, note):
        # row_values[0] is datetime
        path = SNAPSHOT_XLSX
        if os.path.exists(path):
            wb = load_workbook(path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Snapshots"
            ws.append(fields + ["note"])

        ws.append(row_values + [note])
        # format the new time cell
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
        wb.save(path)
        return os.path.abspath(path)

    def _append_to_csv(self, fields, row_values, note):
        path = SNAPSHOT_CSV
        new_file = not os.path.exists(path)
        with open(path, "a", newline="") as f:
            w = csv.writer(f)
            if new_file:
                w.writerow(fields + ["note"])
            w.writerow(row_values + [note])
        return os.path.abspath(path)

    def snapshot_to_excel(self):
        """Append the latest DB row to Excel/CSV with human-friendly time + note."""
        db = self._ensure_db_path()
        if not db:
            QMessageBox.warning(self, "DB not set", "Please locate your SQLite database to take a snapshot.")
            return
        note = self.note_edit.text().strip()

        try:
            con = sqlite3.connect(db)
            fields, row = self._get_latest_row(con)
        except Exception as e:
            QMessageBox.critical(self, "Snapshot error", str(e))
            return
        finally:
            try: con.close()
            except Exception: pass

        # Convert time to datetime + human string
        dt = _parse_db_time(row[0])
        human_time = dt.strftime("%Y-%m-%d %H:%M:%S")  # popular human format
        # For Excel we write the datetime object (so Excel renders/filters properly)
        row_for_excel = [dt] + list(row[1:])
        # For CSV we write the same human string
        row_for_csv   = [human_time] + list(row[1:])

        try:
            if _HAVE_OPENPYXL:
                out = self._append_to_xlsx(fields, row_for_excel, note)
            else:
                out = self._append_to_csv(fields, row_for_csv, note)
        except Exception as e:
            QMessageBox.critical(self, "Write error", f"Failed to write snapshot:\n{e}")
            return

        QMessageBox.information(self, "Snapshot saved", f"Appended 1 row to:\n{out}")


def run_gui(data_queue=None, target_values=None, control_queue=None):
    app = QApplication(sys.argv)
    window = MainControlWindow(data_queue, target_values, control_queue)
    window.show()
    sys.exit(app.exec())
